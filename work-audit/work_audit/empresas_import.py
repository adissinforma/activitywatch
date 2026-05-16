"""Importación de empresas desde un archivo Excel.

Formato esperado: columnas CODEMPRESA, NOMEMPRESA, CIF. Si la primera fila
es una cabecera reconocible se respeta el orden de sus columnas; si no, se
toman las tres primeras columnas en ese orden.
"""

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

log = logging.getLogger(__name__)

_CABECERAS = {"CODEMPRESA", "NOMEMPRESA", "CIF"}


def _celda(fila: tuple, idx: Optional[int]) -> str:
    if idx is None or idx >= len(fila) or fila[idx] is None:
        return ""
    return str(fila[idx]).strip()


def import_excel(path) -> List[Dict[str, Any]]:
    """Lee un Excel y devuelve [{CODEMPRESA, NOMEMPRESA, CIF}, ...].

    Las filas sin CODEMPRESA se ignoran.
    """
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(Path(path)), read_only=True, data_only=True)
    try:
        ws = wb.active
        filas = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not filas:
        return []

    cabecera = [
        str(c).strip().upper() if c is not None else "" for c in filas[0]
    ]
    if _CABECERAS & set(cabecera):
        idx_cod = cabecera.index("CODEMPRESA") if "CODEMPRESA" in cabecera else 0
        idx_nom = cabecera.index("NOMEMPRESA") if "NOMEMPRESA" in cabecera else 1
        idx_cif = cabecera.index("CIF") if "CIF" in cabecera else 2
        datos = filas[1:]
    else:
        idx_cod, idx_nom, idx_cif = 0, 1, 2
        datos = filas

    empresas: List[Dict[str, Any]] = []
    for fila in datos:
        cod = _celda(fila, idx_cod)
        if not cod:
            continue
        empresas.append(
            {
                "CODEMPRESA": cod,
                "NOMEMPRESA": _celda(fila, idx_nom),
                "CIF": _celda(fila, idx_cif),
            }
        )
    return empresas
